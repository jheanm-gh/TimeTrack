"""Saving the workbook: atomic writes, and Excel holding the file open."""

from __future__ import annotations

import datetime as _dt
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.export import saver
from app.export.csv_export import CSV_COLUMNS, timesheet_csv_text
from app.export.gather import gather
from app.export.saver import SaveOutcome, save_workbook, timestamped_name
from tests.conftest import sast


@pytest.fixture
def repo_with_workbook(repo, project, tmp_path):
    repo.set_setting("workbook.path", str(tmp_path / "out"))
    repo.add_manual_entry(
        project,
        started_at=sast(2026, 9, 14, 9),
        ended_at=sast(2026, 9, 14, 11, 30),
        description="Slope stability analysis",
    )
    return repo


def lock_the_target(monkeypatch):
    """Simulate Excel holding the file: the rename is refused."""
    real_replace = os.replace

    def refuse(src, dst):
        if str(dst).endswith("TimeLog.xlsx"):
            raise PermissionError(13, "The process cannot access the file")
        return real_replace(src, dst)

    monkeypatch.setattr(saver.os, "replace", refuse)


class TestNormalSave:
    def test_it_writes_a_workbook_and_a_csv(self, repo_with_workbook):
        outcome = save_workbook(repo_with_workbook)
        assert outcome.saved
        assert outcome.path.name == "TimeLog.xlsx"
        assert outcome.path.exists()
        assert outcome.csv_path.exists()
        assert "Timesheet" in load_workbook(outcome.path).sheetnames

    def test_saving_twice_replaces_the_file(self, repo_with_workbook, project):
        first = save_workbook(repo_with_workbook)
        rows_before = load_workbook(first.path)["Timesheet"].max_row

        repo_with_workbook.add_manual_entry(
            project, started_at=sast(2026, 9, 16, 9), ended_at=sast(2026, 9, 16, 11)
        )
        second = save_workbook(repo_with_workbook)
        assert second.path == first.path
        assert load_workbook(second.path)["Timesheet"].max_row == rows_before + 1

    def test_no_temporary_files_are_left_behind(self, repo_with_workbook):
        outcome = save_workbook(repo_with_workbook)
        leftovers = [p.name for p in outcome.path.parent.glob("*.tmp*")]
        assert leftovers == []

    def test_the_status_message_names_the_time(self, repo_with_workbook):
        assert save_workbook(repo_with_workbook).message.startswith("Last saved ")


class TestFileLockedByExcel:
    def test_an_autosave_gives_up_quietly_and_says_it_will_retry(
        self, repo_with_workbook, monkeypatch
    ):
        lock_the_target(monkeypatch)
        outcome = save_workbook(repo_with_workbook, manual=False)
        assert not outcome.saved
        assert outcome.locked
        assert "will retry" in outcome.message

    def test_an_autosave_leaves_no_timestamped_copies(
        self, repo_with_workbook, monkeypatch, tmp_path
    ):
        """This is how an export folder becomes unusable."""
        lock_the_target(monkeypatch)
        for _ in range(4):
            save_workbook(repo_with_workbook, manual=False)
        folder = tmp_path / "out"
        assert list(folder.glob("TimeLog_*.xlsx")) == []
        assert list(folder.glob("*.tmp*")) == []

    def test_a_manual_save_falls_back_to_a_timestamped_file(
        self, repo_with_workbook, monkeypatch
    ):
        lock_the_target(monkeypatch)
        outcome = save_workbook(repo_with_workbook, manual=True)
        assert outcome.saved
        assert outcome.fell_back
        assert outcome.locked
        assert outcome.path.name.startswith("TimeLog_")
        assert outcome.path.exists()

    def test_the_fallback_message_is_plain_english(self, repo_with_workbook, monkeypatch):
        lock_the_target(monkeypatch)
        message = save_workbook(repo_with_workbook, manual=True).message
        assert "open in Excel" in message
        assert "saved as TimeLog_" in message

    def test_a_manual_save_never_fails_silently(self, repo_with_workbook, monkeypatch):
        lock_the_target(monkeypatch)
        outcome = save_workbook(repo_with_workbook, manual=True)
        assert outcome.needs_telling

    def test_a_quiet_autosave_skip_does_not_interrupt(self, repo_with_workbook, monkeypatch):
        lock_the_target(monkeypatch)
        assert not save_workbook(repo_with_workbook, manual=False).needs_telling


class TestAtomicity:
    def test_a_failed_write_leaves_the_previous_file_intact(
        self, repo_with_workbook, monkeypatch
    ):
        good = save_workbook(repo_with_workbook)
        original = good.path.read_bytes()

        def explode(self, *args, **kwargs):
            raise OSError("disk full")

        monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", explode)
        outcome = save_workbook(repo_with_workbook)

        assert not outcome.saved
        assert not outcome.locked  # a real fault, not Excel being open
        assert good.path.read_bytes() == original

    def test_a_real_failure_is_reported_rather_than_mistaken_for_a_lock(
        self, repo_with_workbook, monkeypatch
    ):
        def explode(src, dst):
            raise OSError(28, "No space left on device")

        monkeypatch.setattr(saver.os, "replace", explode)
        outcome = save_workbook(repo_with_workbook, manual=True)
        assert not outcome.saved
        assert not outcome.locked
        assert "No space left" in outcome.message


class TestNaming:
    def test_the_timestamped_name_reads_clearly(self):
        result = timestamped_name(Path("/x/TimeLog.xlsx"), _dt.datetime(2026, 9, 14, 14, 32))
        assert result.name == "TimeLog_2026-09-14_1432.xlsx"


class TestCsv:
    def test_the_column_order_is_stable(self, repo_with_workbook):
        rows = gather(repo_with_workbook, today=_dt.date(2026, 9, 14)).timesheet
        text = timesheet_csv_text(rows)
        header = text.splitlines()[0]
        assert header.split(",")[:3] == ["Project", "Date", "Day"]
        assert len(header.split(",")) == len(CSV_COLUMNS)

    def test_dates_are_written_in_iso_form(self, repo_with_workbook):
        rows = gather(repo_with_workbook, today=_dt.date(2026, 9, 14)).timesheet
        assert "2026-09-14" in timesheet_csv_text(rows).splitlines()[1]

    def test_the_csv_sits_beside_the_workbook(self, repo_with_workbook):
        outcome = save_workbook(repo_with_workbook)
        assert outcome.csv_path.parent == outcome.path.parent
        assert outcome.csv_path.name == "TimeLog.csv"

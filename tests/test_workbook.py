"""The spreadsheet itself, checked by reading the saved file back.

The brief is specific about the file: real dates and real numbers, frozen
headers, an autofilter on every tabular sheet, and no merged cells anywhere
in a data region. All of those are verified here against a genuine .xlsx.
"""

from __future__ import annotations

import datetime as _dt
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.core.models import EntryKind, TravelDetail
from app.export.workbook import build_from_repository
from tests.conftest import sast

TODAY = _dt.date(2026, 9, 14)


@pytest.fixture
def populated(repo, project):
    """A small but representative database."""
    other = repo.add_project("Rustenburg Slimes Dam")
    repo.add_manual_entry(
        project,
        started_at=sast(2026, 9, 14, 8, 15),
        ended_at=sast(2026, 9, 14, 10, 30),
        description="Slope stability analysis",
    )
    repo.add_manual_entry(
        project,
        started_at=sast(2026, 9, 14, 13),
        ended_at=sast(2026, 9, 14, 16, 45),
        description="Design report",
    )
    repo.add_manual_entry(
        project,
        EntryKind.SOFTWARE,
        started_at=sast(2026, 9, 14, 20),
        ended_at=sast(2026, 9, 14, 23, 30),
        software_name="PLAXIS 3D",
    )
    repo.add_manual_entry(
        other, started_at=sast(2026, 9, 15, 9), ended_at=sast(2026, 9, 15, 12)
    )
    repo.log_travel(
        other,
        on_date=_dt.date(2026, 9, 15),
        travel=TravelDetail(
            odo_start=Decimal("104200"),
            odo_end=Decimal("104442"),
            trip_to="Rustenburg No. 4",
            trip_purpose="Site inspection",
        ),
    )
    repo.set_ticked(project, _dt.date(2026, 9, 14), True)
    return repo


@pytest.fixture
def saved(populated, tmp_path):
    """Build the workbook, save it, and read it back from disk."""
    workbook, data = build_from_repository(populated, today=TODAY)
    path = tmp_path / "TimeLog.xlsx"
    workbook.save(path)
    return load_workbook(path), data, path


class TestStructure:
    def test_the_six_sheets_are_present_and_in_order(self, saved):
        workbook, _data, _path = saved
        names = workbook.sheetnames
        assert names[0] == "Timesheet"
        for expected in ("Raw Log", "Travel", "Projects", "Summary", "Gaps"):
            assert expected in names
        assert names.index("Raw Log") < names.index("Travel") < names.index("Summary")

    def test_per_project_tabs_are_generated(self, saved):
        workbook, _data, _path = saved
        tabs = [name for name in workbook.sheetnames if name.startswith("TS ")]
        assert len(tabs) == 2
        assert any("Kloof" in name for name in tabs)

    def test_per_project_tabs_can_be_switched_off(self, populated, tmp_path):
        populated.set_setting("workbook.per_project_sheets", "0")
        workbook, _data = build_from_repository(populated, today=TODAY)
        assert not [name for name in workbook.sheetnames if name.startswith("TS ")]

    def test_the_number_of_project_tabs_is_capped(self, repo, tmp_path):
        repo.set_setting("workbook.max_project_sheets", "3")
        for index in range(6):
            project_id = repo.add_project(f"Project {index}")
            repo.add_manual_entry(
                project_id, started_at=sast(2026, 9, 14, 9), ended_at=sast(2026, 9, 14, 10)
            )
        workbook, _data = build_from_repository(repo, today=TODAY)
        assert len([n for n in workbook.sheetnames if n.startswith("TS ")]) == 3

    def test_no_merged_cells_anywhere(self, saved):
        """Merged cells break filtering, sorting and copying."""
        workbook, _data, _path = saved
        for sheet in workbook.worksheets:
            assert list(sheet.merged_cells.ranges) == [], sheet.title

    def test_every_tabular_sheet_has_a_frozen_header_and_a_filter(self, saved):
        workbook, _data, _path = saved
        for title in ("Timesheet", "Raw Log", "Travel", "Projects", "Summary", "Gaps"):
            sheet = workbook[title]
            assert sheet.auto_filter.ref, f"{title} has no autofilter"
            assert sheet.freeze_panes, f"{title} has no frozen header"

    def test_columns_have_widths_set(self, saved):
        workbook, _data, _path = saved
        sheet = workbook["Timesheet"]
        assert sheet.column_dimensions["A"].width > 10
        assert sheet.column_dimensions["M"].width > 20  # Description


class TestRealValues:
    def test_dates_are_excel_dates_not_text(self, saved):
        workbook, _data, _path = saved
        cell = workbook["Timesheet"]["B2"]
        assert isinstance(cell.value, _dt.datetime)
        assert cell.number_format == "dd mmm yyyy"

    def test_times_are_excel_times(self, saved):
        workbook, _data, _path = saved
        assert isinstance(workbook["Timesheet"]["D2"].value, _dt.time)

    def test_hours_are_numbers_not_text(self, saved):
        workbook, _data, _path = saved
        sheet = workbook["Timesheet"]
        for column in ("G", "H", "I", "J"):
            cell = sheet[f"{column}2"]
            assert isinstance(cell.value, (int, float)), f"{column} is {type(cell.value)}"
            assert cell.number_format == "0.00"

    def test_the_billed_column_holds_the_rounded_figure(self, saved):
        workbook, _data, _path = saved
        sheet = workbook["Timesheet"]
        headers = [cell.value for cell in sheet[1]]
        billed = headers.index("Hours (billed)") + 1
        raw = headers.index("Hours (raw)") + 1
        # 2h15m + 3h45m = 6h00m exactly, so billed equals raw here.
        assert sheet.cell(row=2, column=raw).value == 6.0
        assert sheet.cell(row=2, column=billed).value == 6.0

    def test_the_billed_column_is_visually_distinct(self, saved):
        workbook, _data, _path = saved
        sheet = workbook["Timesheet"]
        headers = [cell.value for cell in sheet[1]]
        billed = sheet.cell(row=2, column=headers.index("Hours (billed)") + 1)
        plain = sheet.cell(row=2, column=headers.index("Hours (raw)") + 1)
        assert billed.font.bold
        assert not plain.font.bold
        assert billed.fill.fgColor.rgb != plain.fill.fgColor.rgb

    def test_kilometres_are_numbers(self, saved):
        workbook, _data, _path = saved
        sheet = workbook["Travel"]
        assert isinstance(sheet["D2"].value, (int, float))
        assert sheet["D2"].value == 242.0


class TestTimesheetSheet:
    def test_headers_match_the_brief(self, saved):
        workbook, _data, _path = saved
        assert [cell.value for cell in workbook["Timesheet"][1]] == [
            "Project",
            "Date",
            "Day",
            "First Start",
            "Last End",
            "Sessions",
            "Hours (raw)",
            "Hours (billed)",
            "Software Hours (raw)",
            "Software Hours (billed)",
            "Software Used",
            "Km",
            "Description",
            "Submitted",
        ]

    def test_ticked_dates_are_greyed_by_conditional_formatting(self, saved):
        workbook, _data, _path = saved
        sheet = workbook["Timesheet"]
        rules = [
            rule
            for rng in sheet.conditional_formatting
            for rule in rng.rules
        ]
        assert rules, "no conditional formatting on the Timesheet sheet"
        assert any('="Yes"' in (rule.formula[0] if rule.formula else "") for rule in rules)

    def test_the_submitted_column_shows_the_tick(self, saved):
        workbook, _data, _path = saved
        sheet = workbook["Timesheet"]
        values = {
            (row[0].value, row[13].value)
            for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row)
        }
        assert ("Kloof Tailings Dam", "Yes") in values

    def test_project_tabs_show_the_same_figures_without_the_project_column(self, saved):
        """Acceptance checklist 14."""
        workbook, _data, _path = saved
        tab = next(workbook[n] for n in workbook.sheetnames if n.startswith("TS Kloof"))
        assert [cell.value for cell in tab[1]][0] == "Date"
        assert "Project" not in [cell.value for cell in tab[1]]

        main = workbook["Timesheet"]
        main_row = next(
            row for row in main.iter_rows(min_row=2, values_only=True)
            if row[0] == "Kloof Tailings Dam"
        )
        tab_row = next(tab.iter_rows(min_row=2, max_row=2, values_only=True))
        assert main_row[1:] == tab_row  # same values, Project column removed


class TestTravelSheet:
    def test_columns_follow_the_sars_logbook_order(self, saved):
        workbook, _data, _path = saved
        headers = [cell.value for cell in workbook["Travel"][1]]
        assert headers[:6] == [
            "Date",
            "Opening odometer",
            "Closing odometer",
            "Kilometres",
            "Destination",
            "Reason for trip",
        ]

    def test_missing_odometer_cells_are_blank_not_invented(self, repo, project, tmp_path):
        repo.log_travel(
            project,
            on_date=_dt.date(2026, 9, 14),
            travel=TravelDetail(km_travelled=Decimal("240")),
        )
        workbook, _data = build_from_repository(repo, today=TODAY)
        sheet = workbook["Travel"]
        assert sheet["B2"].value is None
        assert sheet["C2"].value is None
        assert sheet["D2"].value == Decimal("240.0")

    def test_totals_appear_below_the_table(self, saved):
        workbook, _data, _path = saved
        sheet = workbook["Travel"]
        labels = [
            sheet.cell(row=row, column=1).value
            for row in range(1, sheet.max_row + 1)
        ]
        assert "Totals" in labels
        assert "By month" in labels
        assert "By project" in labels

    def test_the_totals_are_outside_the_filter_range(self, saved):
        """Filtering the table must not hide the totals underneath it."""
        workbook, _data, _path = saved
        sheet = workbook["Travel"]
        filter_last_row = int(sheet.auto_filter.ref.split(":")[1].lstrip("ABCDEFGHIJ"))
        totals_row = next(
            row
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row=row, column=1).value == "Totals"
        )
        assert totals_row > filter_last_row


class TestSummarySheet:
    def test_the_rounding_difference_is_shown(self, saved):
        workbook, _data, _path = saved
        headers = [cell.value for cell in workbook["Summary"][1]]
        assert "Rounding difference" in headers

    def test_a_grand_total_row_is_written_clear_of_the_table(self, saved):
        workbook, data, _path = saved
        sheet = workbook["Summary"]
        total_row = 1 + len(data.summary) + 2
        assert sheet.cell(row=total_row, column=1).value == "ALL PROJECTS"
        assert sheet.cell(row=total_row, column=1).font.bold


class TestEmptyWorkbook:
    def test_a_brand_new_database_still_produces_a_valid_file(self, repo, tmp_path):
        """First launch: no entries at all, and the file must still open."""
        workbook, _data = build_from_repository(repo, today=TODAY)
        path = tmp_path / "Empty.xlsx"
        workbook.save(path)

        reopened = load_workbook(path)
        assert "Timesheet" in reopened.sheetnames
        sheet = reopened["Timesheet"]
        assert [cell.value for cell in sheet[1]][0] == "Project"
        # The filter is there ready for the first entry.
        assert sheet.auto_filter.ref == "A1:N1"

"""Building the workbook.

Six sheets, plus optional per-project tabs. Every tabular sheet gets a
frozen header, an autofilter, real dates and real numbers, sensible column
widths and banded rows - and nowhere a merged cell inside a data region.
"""

from __future__ import annotations

import datetime as _dt
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.export.gather import WorkbookData, gather
from app.export.styles import (
    FMT_DATE,
    FMT_HOURS,
    FMT_KM,
    FMT_ODOMETER,
    FMT_TIME,
    TOTALS_FILL,
    TOTALS_FONT,
    Column,
    StyleRegistry,
    finish_table,
    grey_out_when,
    sanitise_sheet_name,
    write_header,
    write_rows,
    write_title,
    write_totals_block,
)
from app.db.repository import Repository

PROJECT_SHEET_PREFIX = "TS "

# -- column definitions ----------------------------------------------------

TIMESHEET_COLUMNS = [
    Column("Project", "project", 34),
    Column("Date", "date", 13, FMT_DATE),
    Column("Day", "day", 6, align="centre"),
    Column("First Start", "first_start", 11, FMT_TIME, align="centre"),
    Column("Last End", "last_end", 11, FMT_TIME, align="centre"),
    Column("Sessions", "sessions", 30),
    Column("Hours (raw)", "hours_raw", 12, FMT_HOURS, align="right"),
    Column("Hours (billed)", "hours_billed", 15, FMT_HOURS, align="right", emphasis=True),
    Column("Software Hours (raw)", "software_raw", 14, FMT_HOURS, align="right"),
    Column(
        "Software Hours (billed)",
        "software_billed",
        16,
        FMT_HOURS,
        align="right",
        emphasis=True,
    ),
    Column("Software Used", "software_used", 22),
    Column("Km", "km", 9, FMT_KM, align="right"),
    Column("Description", "description", 60),
    Column("Submitted", "submitted", 11, align="centre"),
]

#: The per-project tabs show the same figures without the Project column.
PROJECT_TAB_COLUMNS = [column for column in TIMESHEET_COLUMNS if column.attr != "project"]

RAW_LOG_COLUMNS = [
    Column("Date", "date", 13, FMT_DATE),
    Column("Project", "project", 30),
    Column("Task", "task", 24),
    Column("Type", "kind", 10, align="centre"),
    Column("Start", "start", 9, FMT_TIME, align="centre"),
    Column("End", "end", 9, FMT_TIME, align="centre"),
    Column("Duration", "duration_hm", 10, align="right"),
    Column("Hours", "duration_hours", 9, FMT_HOURS, align="right"),
    Column("Software", "software_name", 18),
    Column("Km", "km", 9, FMT_KM, align="right"),
    Column("Odo start", "odo_start", 11, FMT_ODOMETER, align="right"),
    Column("Odo end", "odo_end", 11, FMT_ODOMETER, align="right"),
    Column("From", "trip_from", 20),
    Column("To", "trip_to", 22),
    Column("Purpose", "trip_purpose", 22),
    Column("Description", "description", 50),
    Column("Source", "source", 10, align="centre"),
    Column("Edited", "edited", 8, align="centre"),
    Column("Entry id", "entry_id", 9, align="right"),
]

#: Ordered the way a SARS travel logbook wants it: date, opening reading,
#: closing reading, kilometres, destination, reason. The extra columns the
#: timesheet cares about follow afterwards so one sheet serves both.
TRAVEL_COLUMNS = [
    Column("Date", "date", 13, FMT_DATE),
    Column("Opening odometer", "odo_start", 17, FMT_ODOMETER, align="right"),
    Column("Closing odometer", "odo_end", 17, FMT_ODOMETER, align="right"),
    Column("Kilometres", "km", 12, FMT_KM, align="right", emphasis=True),
    Column("Destination", "destination", 30),
    Column("Reason for trip", "reason", 34),
    Column("From", "trip_from", 22),
    Column("Project", "project", 32),
    Column("Entry id", "entry_id", 9, align="right"),
]

PROJECT_COLUMNS = [
    Column("Project", "name", 36),
    Column("Client", "client", 22),
    Column("Code", "code", 12),
    Column("Status", "status", 10, align="centre"),
    Column("Submission day", "submission_day", 17),
    Column("Period", "period_type", 22),
    Column("Next deadline", "next_deadline", 14, FMT_DATE),
    Column("Usual site", "default_site", 26),
    Column("Usual km", "default_km", 10, FMT_KM, align="right"),
    Column("Hours to date", "hours_to_date", 13, FMT_HOURS, align="right"),
    Column("Billed to date", "billed_to_date", 13, FMT_HOURS, align="right", emphasis=True),
    Column("Km to date", "km_to_date", 12, FMT_KM, align="right"),
    Column("Last submitted period", "last_submitted", 26),
]

SUMMARY_COLUMNS = [
    Column("Project", "project", 36),
    Column("Month", "month_label", 16),
    Column("Hours (raw)", "hours_raw", 12, FMT_HOURS, align="right"),
    Column("Hours (billed)", "hours_billed", 14, FMT_HOURS, align="right", emphasis=True),
    Column("Rounding difference", "rounding_difference", 18, FMT_HOURS, align="right"),
    Column("Software (raw)", "software_raw", 14, FMT_HOURS, align="right"),
    Column("Software (billed)", "software_billed", 15, FMT_HOURS, align="right"),
    Column("Km", "km", 10, FMT_KM, align="right"),
    Column("Submitted", "submitted", 11, align="centre"),
]

GAP_COLUMNS = [
    Column("Date", "date", 14, FMT_DATE),
    Column("Day", "day", 8, align="centre"),
]


def build_workbook(data: WorkbookData, per_project_sheets: bool = True, max_project_sheets: int = 40) -> Workbook:
    """Assemble the whole workbook from already-gathered data."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    registry = StyleRegistry(workbook)

    _write_timesheet(workbook, data, registry)
    if per_project_sheets:
        _write_project_tabs(workbook, data, max_project_sheets, registry)
    _write_raw_log(workbook, data, registry)
    _write_travel(workbook, data, registry)
    _write_projects(workbook, data, registry)
    _write_summary(workbook, data, registry)
    _write_gaps(workbook, data, registry)

    workbook.properties.title = "TimeTrack timesheet"
    workbook.properties.creator = "TimeTrack"
    return workbook


def _simple_sheet(
    workbook: Workbook,
    title: str,
    columns: list[Column],
    rows: list,
    registry: StyleRegistry,
) -> Worksheet:
    sheet = workbook.create_sheet(title)
    write_header(sheet, columns)
    last_row = write_rows(sheet, columns, rows, registry=registry)
    finish_table(sheet, columns, header_row=1, last_row=last_row)
    return sheet


def _write_timesheet(
    workbook: Workbook, data: WorkbookData, registry: StyleRegistry
) -> None:
    """Sheet 1 - every project on one sheet, one row per project per date."""
    sheet = _simple_sheet(
        workbook, "Timesheet", TIMESHEET_COLUMNS, data.timesheet, registry
    )
    grey_out_when(
        sheet,
        TIMESHEET_COLUMNS,
        "submitted",
        header_row=1,
        last_row=1 + len(data.timesheet),
    )


def _write_project_tabs(
    workbook: Workbook, data: WorkbookData, limit: int, registry: StyleRegistry
) -> None:
    """One tab per project, same figures minus the Project column.

    Capped: a workbook with a hundred tabs is slower to open and harder to
    navigate than the filterable sheet it duplicates.
    """
    by_project: dict[str, list] = {}
    for row in data.timesheet:
        by_project.setdefault(row.project, []).append(row)

    used: set[str] = {sheet.title for sheet in workbook.worksheets}
    for project, rows in sorted(by_project.items(), key=lambda kv: kv[0].lower())[:limit]:
        title = sanitise_sheet_name(project, PROJECT_SHEET_PREFIX, used)
        sheet = workbook.create_sheet(title)
        write_header(sheet, PROJECT_TAB_COLUMNS)
        last_row = write_rows(sheet, PROJECT_TAB_COLUMNS, rows, registry=registry)
        finish_table(sheet, PROJECT_TAB_COLUMNS, header_row=1, last_row=last_row)
        grey_out_when(
            sheet, PROJECT_TAB_COLUMNS, "submitted", header_row=1, last_row=last_row
        )


def _write_raw_log(
    workbook: Workbook, data: WorkbookData, registry: StyleRegistry
) -> None:
    """Sheet 2 - every individual entry, the audit trail."""
    _simple_sheet(workbook, "Raw Log", RAW_LOG_COLUMNS, data.raw_log, registry)


def _write_travel(
    workbook: Workbook, data: WorkbookData, registry: StyleRegistry
) -> None:
    """Sheet 3 - the kilometre log, in SARS logbook order."""
    sheet = _simple_sheet(workbook, "Travel", TRAVEL_COLUMNS, data.travel, registry)
    last_row = 1 + len(data.travel)

    # Totals go *below* the table with a gap: Excel's autofilter only hides
    # rows inside its own range, so a filtered view keeps these visible.
    cursor = last_row + 3
    totals = data.travel_totals
    cursor = write_totals_block(
        sheet,
        cursor,
        "Totals",
        [
            (f"Year to date ({data.today.year})", totals.year_to_date),
            ("All time", totals.all_time),
        ],
        FMT_KM,
    )
    cursor += 1
    cursor = write_totals_block(
        sheet, cursor, "By month", totals.by_month, FMT_KM
    )
    cursor += 1
    write_totals_block(sheet, cursor, "By project", totals.by_project, FMT_KM)


def _write_projects(
    workbook: Workbook, data: WorkbookData, registry: StyleRegistry
) -> None:
    """Sheet 4 - the project register."""
    _simple_sheet(workbook, "Projects", PROJECT_COLUMNS, data.projects, registry)


def _write_summary(
    workbook: Workbook, data: WorkbookData, registry: StyleRegistry
) -> None:
    """Sheet 5 - per project per month, with the rounding difference."""
    sheet = _simple_sheet(workbook, "Summary", SUMMARY_COLUMNS, data.summary, registry)
    last_row = 1 + len(data.summary)

    if data.grand_total is not None:
        # Two rows clear of the table: Excel's autofilter only hides rows
        # inside its own range, so the grand total survives any filtering.
        total_row = last_row + 2
        write_rows(
            sheet,
            SUMMARY_COLUMNS,
            [data.grand_total],
            first_row=total_row,
            registry=registry,
        )
        for index in range(1, len(SUMMARY_COLUMNS) + 1):
            cell = sheet.cell(row=total_row, column=index)
            cell.font = TOTALS_FONT
            cell.fill = TOTALS_FILL

        note_row = total_row + 2
        write_title(
            sheet,
            "About the rounding difference",
            row=note_row,
            note=(
                "Each day's total is rounded up to the next quarter hour, once "
                "per project per day. The rounding difference column is how "
                "many hours that adds over the month."
            ),
        )


def _write_gaps(
    workbook: Workbook, data: WorkbookData, registry: StyleRegistry
) -> None:
    """Sheet 6 - weekdays with nothing recorded, still to be reconstructed."""
    sheet = workbook.create_sheet("Gaps")
    next_row = write_title(
        sheet,
        "Weekdays with no time recorded",
        row=1,
        note=(
            "Working days in the recent past that have no hours against any "
            "project. Today is not listed - the day is not over yet."
        ),
    )
    header_row = next_row + 1
    write_header(sheet, GAP_COLUMNS, row=header_row)
    last_row = write_rows(
        sheet, GAP_COLUMNS, data.gaps, first_row=header_row + 1, registry=registry
    )
    finish_table(sheet, GAP_COLUMNS, header_row=header_row, last_row=last_row)


def build_from_repository(
    repo: Repository, today: _dt.date | None = None
) -> tuple[Workbook, WorkbookData]:
    """Gather and build in one step, honouring the workbook settings."""
    data = gather(repo, today=today)
    workbook = build_workbook(
        data,
        per_project_sheets=repo.get_bool("workbook.per_project_sheets", True),
        max_project_sheets=repo.get_int("workbook.max_project_sheets", 40),
    )
    return workbook, data
